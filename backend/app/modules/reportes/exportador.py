"""Generación de reportes exportables en PDF y Excel (CU13 — Consultar reportes y dashboards)."""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.modules.reportes import service
from app.shared.core.tiempo import ZONA_HORARIA_LOCAL

COLOR_HEADER = "1F2937"
COLOR_ZEBRA = "F3F4F6"
COLOR_MUTED = "6B7280"


def _construir_datos(db: Session, sucursal_id: int | None, es_admin: bool) -> dict:
    return {
        "resumen": service.resumen_general(db, sucursal_id=sucursal_id),
        "por_estado": service.reservas_por_estado(db, sucursal_id=sucursal_id),
        "top_productos": service.productos_mas_reservados(db, sucursal_id=sucursal_id, limit=15),
        "por_dia": service.reservas_por_dia(db, sucursal_id=sucursal_id, dias=30),
        "por_sucursal": service.reservas_por_sucursal(db) if es_admin and sucursal_id is None else [],
    }


def generar_excel(db: Session, sucursal_id: int | None, nombre_sucursal: str, es_admin: bool) -> bytes:
    datos = _construir_datos(db, sucursal_id, es_admin)
    resumen = datos["resumen"]

    wb = Workbook()
    header_fill = PatternFill(start_color=COLOR_HEADER, end_color=COLOR_HEADER, fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(bold=True, size=14, color=COLOR_HEADER)

    def _hoja_tabla(ws, encabezados: list[str], filas: list[list], titulo: str):
        ws["A1"] = titulo
        ws["A1"].font = title_font
        ws.append([])
        ws.append(encabezados)
        fila_encabezado = ws.max_row
        for col in range(1, len(encabezados) + 1):
            celda = ws.cell(row=fila_encabezado, column=col)
            celda.fill = header_fill
            celda.font = header_font
            celda.alignment = Alignment(horizontal="center")
        for fila in filas:
            ws.append(fila)
        for col in range(1, len(encabezados) + 1):
            anchos = [len(str(encabezados[col - 1]))] + [len(str(f[col - 1])) for f in filas]
            ws.column_dimensions[get_column_letter(col)].width = min(max(anchos) + 4, 42)

    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    _hoja_tabla(
        ws_resumen,
        ["Indicador", "Valor"],
        [
            ["Reservas totales", resumen["total_reservas"]],
            ["Unidades reservadas", resumen["total_unidades_reservadas"]],
            ["Inventario disponible", resumen["inventario_disponible"]],
            ["Inventario reservado", resumen["inventario_reservado"]],
            ["Productos con stock bajo", resumen["productos_stock_bajo"]],
            ["Productos agotados", resumen["productos_agotados"]],
        ],
        f"Resumen de reportes — {nombre_sucursal}",
    )

    _hoja_tabla(
        wb.create_sheet("Reservas por estado"),
        ["Estado", "Cantidad"],
        [[e["estado"].capitalize(), e["cantidad"]] for e in datos["por_estado"]] or [["Sin datos", 0]],
        "Reservas por estado",
    )

    _hoja_tabla(
        wb.create_sheet("Productos mas reservados"),
        ["Producto", "Categoría", "Unidades", "Reservas"],
        [
            [p["nombre_producto"], p["nombre_categoria"] or "Sin categoría", p["total_unidades"], p["total_reservas"]]
            for p in datos["top_productos"]
        ]
        or [["Sin datos", "-", 0, 0]],
        "Productos más reservados",
    )

    _hoja_tabla(
        wb.create_sheet("Reservas por dia"),
        ["Fecha", "Cantidad"],
        [[d["fecha"], d["cantidad"]] for d in datos["por_dia"]] or [["Sin datos", 0]],
        "Reservas por día (últimos 30 días)",
    )

    if datos["por_sucursal"]:
        _hoja_tabla(
            wb.create_sheet("Comparativa sucursales"),
            ["Sucursal", "Reservas", "Unidades reservadas"],
            [[s["nombre_sucursal"], s["cantidad_reservas"], s["total_unidades"]] for s in datos["por_sucursal"]],
            "Comparativa entre sucursales",
        )

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generar_pdf(db: Session, sucursal_id: int | None, nombre_sucursal: str, es_admin: bool) -> bytes:
    datos = _construir_datos(db, sucursal_id, es_admin)
    resumen = datos["resumen"]

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        title="Reporte SmartLook-AI",
    )

    styles = getSampleStyleSheet()
    titulo_style = ParagraphStyle("TituloReporte", parent=styles["Title"], textColor=colors.HexColor(f"#{COLOR_HEADER}"))
    subtitulo_style = ParagraphStyle(
        "Subtitulo", parent=styles["Normal"], textColor=colors.HexColor(f"#{COLOR_MUTED}"), spaceAfter=6
    )
    seccion_style = ParagraphStyle(
        "Seccion", parent=styles["Heading2"], textColor=colors.HexColor(f"#{COLOR_HEADER}"), spaceBefore=16, spaceAfter=6
    )

    def _tabla(encabezados: list[str], filas: list[list[str]]) -> Table:
        tabla = Table([encabezados] + filas, repeatRows=1)
        tabla.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{COLOR_HEADER}")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{COLOR_ZEBRA}")]),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        return tabla

    elementos = [
        Paragraph("SmartLook-AI — Reporte de reservas e inventario", titulo_style),
        Paragraph(
            f"{nombre_sucursal} &bull; Generado el "
            f"{datetime.now(ZONA_HORARIA_LOCAL).strftime('%d/%m/%Y %H:%M')}",
            subtitulo_style,
        ),
        Spacer(1, 0.3 * cm),
        Paragraph("Resumen general", seccion_style),
        _tabla(
            ["Indicador", "Valor"],
            [
                ["Reservas totales", str(resumen["total_reservas"])],
                ["Unidades reservadas", str(resumen["total_unidades_reservadas"])],
                ["Inventario disponible", str(resumen["inventario_disponible"])],
                ["Inventario reservado", str(resumen["inventario_reservado"])],
                ["Productos con stock bajo", str(resumen["productos_stock_bajo"])],
                ["Productos agotados", str(resumen["productos_agotados"])],
            ],
        ),
    ]

    if datos["por_estado"]:
        elementos.append(Paragraph("Reservas por estado", seccion_style))
        elementos.append(
            _tabla(
                ["Estado", "Cantidad"],
                [[e["estado"].capitalize(), str(e["cantidad"])] for e in datos["por_estado"]],
            )
        )

    if datos["por_sucursal"]:
        elementos.append(Paragraph("Comparativa entre sucursales", seccion_style))
        elementos.append(
            _tabla(
                ["Sucursal", "Reservas", "Unidades reservadas"],
                [[s["nombre_sucursal"], str(s["cantidad_reservas"]), str(s["total_unidades"])] for s in datos["por_sucursal"]],
            )
        )

    if datos["top_productos"]:
        elementos.append(Paragraph("Productos más reservados", seccion_style))
        elementos.append(
            _tabla(
                ["Producto", "Categoría", "Unidades", "Reservas"],
                [
                    [
                        p["nombre_producto"],
                        p["nombre_categoria"] or "Sin categoría",
                        str(p["total_unidades"]),
                        str(p["total_reservas"]),
                    ]
                    for p in datos["top_productos"]
                ],
            )
        )

    if datos["por_dia"]:
        elementos.append(Paragraph("Reservas por día (últimos 30 días)", seccion_style))
        elementos.append(
            _tabla(
                ["Fecha", "Cantidad"],
                [[d["fecha"], str(d["cantidad"])] for d in datos["por_dia"]],
            )
        )

    doc.build(elementos)
    return buffer.getvalue()
